"""Convert the corrected legacy faturamento worksheet to the versioned Hub seed.

Developer-only tool: the application and build consume the generated JSON and do
not require XLSX, Python, or openpyxl at runtime.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


SEED_VERSION = 1
SEED_SCHEMA = "billing_seed_v1"
SOURCE = "legacy_faturamento_xlsx"
MONTH_NAMES = {
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12,
}


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def iso_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def number(value: object) -> int | float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    result = float(value)
    return int(result) if result.is_integer() else result


def norm(value: object) -> str:
    text = unicodedata.normalize("NFD", clean_text(value).lower())
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def status_for(*values: object) -> str:
    for value in values:
        token = norm(value)
        if token in {"pago", "paga", "paid", "recebido", "recebida"}:
            return "paid"
        if token in {"pendente", "pending"}:
            return "pending"
        if token in {"parcial", "parcialmente pago", "partially paid"}:
            return "partially_paid"
    return "unknown"


def stable_id(prefix: str, fingerprint: str, occurrence: int = 1) -> str:
    digest = hashlib.sha256(f"{fingerprint}\0{occurrence}".encode("utf-8")).hexdigest()[:24]
    return f"{prefix}-{digest}"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Workbook original corrigido")
    parser.add_argument("--output", required=True, type=Path, help="Destino do JSON de seed")
    args = parser.parse_args()

    workbook = load_workbook(args.input, data_only=True, read_only=True)
    if "fatur." not in workbook.sheetnames:
        raise SystemExit("A aba obrigatória 'fatur.' não foi encontrada.")
    sheet = workbook["fatur."]
    raw_sales: list[dict] = []
    raw_refunds: list[dict] = []
    ambiguous_rows: list[int] = []
    ignored = Counter()
    monthly = defaultdict(lambda: {"sales": 0, "brl": 0.0, "usd": 0.0})
    monthly_source: dict[str, float] = {}
    status_counts = Counter()
    missing_usd = 0
    sale_brl_total = 0.0
    sale_usd_total = 0.0
    refund_brl_total = 0.0
    refund_usd_missing = 0
    dated_payment_count = 0
    paid_without_date_count = 0
    unknown_without_refund_count = 0

    for row_number in range(1, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, 14)]
        sequence, sale_date_raw, brl_raw, usd_raw, platform, commission_type, product, account, column_i, column_j = values[:10]
        sale_date = iso_date(sale_date_raw)
        brl = number(brl_raw)
        usd = number(usd_raw)
        candidate = isinstance(sequence, (int, float)) and not isinstance(sequence, bool)
        if candidate and sale_date and brl is not None:
            if not clean_text(platform) or not clean_text(product):
                ambiguous_rows.append(row_number)
                continue
            refund_brl = number(column_i)
            refund_date = iso_date(column_j) if refund_brl is not None else None
            payment_date = iso_date(column_i) or (iso_date(column_j) if refund_brl is None else None)
            payment_status = status_for(column_i, column_j)
            if payment_status == "unknown" and payment_date:
                payment_status = "paid"
            notes = []
            for value in (column_i, column_j):
                if value is None or isinstance(value, (int, float, datetime, date)):
                    continue
                token = norm(value)
                if token not in {"pago", "paga", "paid", "recebido", "recebida", "pendente", "pending", "parcial", "parcialmente pago", "partially paid"}:
                    notes.append(clean_text(value))
            fingerprint_fields = [sale_date, clean_text(platform), clean_text(commission_type), clean_text(product), clean_text(account), brl, usd, payment_status, payment_date, refund_brl, refund_date, " · ".join(dict.fromkeys(notes))]
            fingerprint = json.dumps(fingerprint_fields, ensure_ascii=False, separators=(",", ":"))
            raw_sales.append({
                "row_number": row_number,
                "fingerprint": fingerprint,
                "sale_date": sale_date,
                "platform": clean_text(platform),
                "commission_type": clean_text(commission_type),
                "product": clean_text(product),
                "account": clean_text(account),
                "value_brl": brl,
                "value_usd": usd,
                "payment_status": payment_status,
                "payment_date": payment_date if payment_status == "paid" else None,
                "refund_brl": refund_brl,
                "refund_date": refund_date,
                "notes": " · ".join(dict.fromkeys(notes)),
                "status_conflict": payment_status == "pending" and payment_date is not None,
            })
            status_counts[payment_status] += 1
            if payment_status == "unknown" and refund_brl is None:
                unknown_without_refund_count += 1
            sale_brl_total += brl
            if usd is None:
                missing_usd += 1
            else:
                sale_usd_total += usd
            month_key = sale_date[:7]
            monthly[month_key]["sales"] += 1
            monthly[month_key]["brl"] += brl
            if usd is not None:
                monthly[month_key]["usd"] += usd
            if payment_status == "paid" and payment_date:
                dated_payment_count += 1
            elif payment_status == "paid":
                paid_without_date_count += 1
            if refund_brl is not None:
                raw_refunds.append({"row_number": row_number, "fingerprint": fingerprint, "sale_date": sale_date, "value_brl": refund_brl, "effective_date": refund_date, "notes": "Reembolso registrado na fonte histórica."})
                refund_brl_total += refund_brl
                refund_usd_missing += 1
            if payment_status == "pending" and payment_date:
                ambiguous_rows.append(row_number)
            continue

        nonempty = any(value is not None and value != "" for value in values)
        if not nonempty:
            continue
        row_text = " ".join(norm(value) for value in values if value is not None)
        if any(token in row_text for token in ("data", "valor comiss", "plataforma", "reembolso", "pagamento", "faturamento (r$)")):
            ignored["headers"] += 1
        elif "total" in row_text or "pendente" in row_text:
            ignored["totals_or_status_summaries"] += 1
        else:
            ignored["titles_or_summary_area"] += 1

    occurrences = Counter()
    sales = []
    row_to_sale_id: dict[int, str] = {}
    for row in raw_sales:
        occurrences[row["fingerprint"]] += 1
        sale_id = stable_id("legacy-sale", row["fingerprint"], occurrences[row["fingerprint"]])
        row_to_sale_id[row["row_number"]] = sale_id
        sales.append({
            "sale_id": sale_id,
            "sale_date": row["sale_date"],
            "platform": row["platform"],
            "product": row["product"],
            "commission_type": row["commission_type"],
            "account": row["account"],
            "value_brl": row["value_brl"],
            "value_usd": row["value_usd"],
            "payment_status": row["payment_status"],
            "observed_payment_status": row["payment_status"],
            "notes": row["notes"],
            "source": SOURCE,
            "source_ref": f"fatur.#{row['row_number']}",
            "external_id": None,
            "active": True,
            "cancelled_at": None,
            "cancellation_reason": "",
            "created_at": "2026-09-24T00:00:00.000Z",
            "updated_at": "2026-09-24T00:00:00.000Z",
        })

    movements = []
    movement_occurrences = Counter()
    for row in raw_sales:
        sale_id = row_to_sale_id[row["row_number"]]
        if row["payment_status"] == "paid" and row["payment_date"]:
            fingerprint = json.dumps([sale_id, "receipt", row["payment_date"], row["value_brl"], row["value_usd"]], separators=(",", ":"))
            movement_occurrences[fingerprint] += 1
            movements.append({
                "movement_id": stable_id("legacy-movement", fingerprint, movement_occurrences[fingerprint]),
                "sale_id": sale_id,
                "type": "receipt",
                "effective_date": row["payment_date"],
                "value_brl": row["value_brl"],
                "value_usd": row["value_usd"],
                "source": SOURCE,
                "notes": "Data de pagamento explícita na fonte histórica.",
                "created_at": "2026-09-24T00:00:00.000Z",
            })
        if row["refund_brl"] is not None:
            fingerprint = json.dumps([sale_id, "refund", row["refund_date"], row["refund_brl"]], separators=(",", ":"))
            movement_occurrences[fingerprint] += 1
            movements.append({
                "movement_id": stable_id("legacy-movement", fingerprint, movement_occurrences[fingerprint]),
                "sale_id": sale_id,
                "type": "refund",
                "effective_date": row["refund_date"],
                "value_brl": row["refund_brl"],
                "value_usd": None,
                "source": SOURCE,
                "notes": row["notes"] or "Reembolso registrado na fonte histórica; valor em US$ não informado.",
                "created_at": "2026-09-24T00:00:00.000Z",
            })

    controls = {}
    for row_number in range(1, sheet.max_row + 1):
        label = clean_text(sheet.cell(row_number, 12).value)
        value = number(sheet.cell(row_number, 13).value)
        match = re.fullmatch(r"([A-Za-zÀ-ÿ]{3})/(\d{2})", label)
        if match and value is not None:
            month = MONTH_NAMES.get(norm(match.group(1))[:3])
            if month:
                year = 2000 + int(match.group(2))
                controls[f"{year:04d}-{month:02d}"] = value
    monthly_validation = []
    for key, values in sorted(monthly.items()):
        control = controls.get(key)
        difference = values["brl"] - control if control is not None else None
        monthly_validation.append({
            "month": key, "sales": values["sales"], "brl_source_lines": values["brl"],
            "usd_source_lines": values["usd"], "brl_monthly_control": control,
            "brl_difference": difference,
        })

    # Any date/status contradiction or incomplete identity-critical row remains for manual review.
    ambiguous_rows = sorted(set(ambiguous_rows))
    if ambiguous_rows:
        raise SystemExit(f"Há linhas ambíguas que precisam de decisão antes de gerar o seed: {ambiguous_rows}")
    for row in raw_sales:
        if not row["sale_date"] or row["value_brl"] is None or not row["product"] or not row["platform"]:
            raise SystemExit(f"Registro incompleto na origem fatura.# {row['row_number']}; seed não gerado.")

    seed = {
        "schema": SEED_SCHEMA,
        "version": SEED_VERSION,
        "source": {"type": SOURCE, "workbook": args.input.name, "sheet": "fatur.", "generated_by": "scripts/generate-billing-seed.py"},
        "validation": {
            "sales": len(sales),
            "payments": dict(status_counts),
            "paid_with_known_date": dated_payment_count,
            "paid_without_known_date": paid_without_date_count,
            "unknown_payment_status_without_refund": unknown_without_refund_count,
            "refund_movements": len(raw_refunds),
            "refund_brl_total": refund_brl_total,
            "refund_usd_missing": refund_usd_missing,
            "usd_missing_sales": missing_usd,
            "ignored_source_rows": dict(ignored),
            "ambiguous_rows": len(ambiguous_rows),
            "gross_totals": {"brl": sale_brl_total, "usd": sale_usd_total},
            "monthly_validation": monthly_validation,
        },
        "sales": sales,
        "movements": movements,
        "audit": [],
        "meta": [{"key": "seed:legacy-faturamento:v1", "version": SEED_VERSION, "source": SOURCE, "sales_total": len(sales), "movements_total": len(movements)}],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(seed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(seed["validation"], ensure_ascii=False, indent=2))
    print(f"seed escrito: {args.output}")


if __name__ == "__main__":
    main()
